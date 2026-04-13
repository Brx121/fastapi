from __future__ import annotations

import os
from typing import Any, Dict, List, Optional

import PyPDF2
from docx import Document
from openpyxl import load_workbook


def parse_pdf(file_path: str) -> str:
    """
    解析 PDF 文档并提取文本
    """
    text = ""
    with open(file_path, 'rb') as f:
        reader = PyPDF2.PdfReader(f)
        for page_num in range(len(reader.pages)):
            page = reader.pages[page_num]
            text += page.extract_text() + "\n"
    return text


def parse_excel(file_path: str) -> str:
    """
    解析 Excel 文档并提取文本
    """
    text = ""
    wb = load_workbook(file_path)
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        text += f"Sheet: {sheet_name}\n"
        for row in sheet.iter_rows(values_only=True):
            row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
            if row_text.strip():
                text += row_text + "\n"
        text += "\n"
    return text


def parse_docx(file_path: str) -> str:
    """
    解析 Word 文档并提取文本
    """
    doc = Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    return text


def parse_document(file_path: str) -> str:
    """
    根据文件扩展名自动选择解析器
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return parse_pdf(file_path)
    elif ext in ['.xlsx', '.xls']:
        return parse_excel(file_path)
    elif ext == '.docx':
        return parse_docx(file_path)
    else:
        # 对于其他文件类型，尝试直接读取文本
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception:
            raise ValueError(f"Unsupported file type: {ext}")
